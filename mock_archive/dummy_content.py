import os
from docx import Document
from openpyxl import Workbook
from fpdf import FPDF

def create_docx(path, text):
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)

def create_xlsx(path, text):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = text
    wb.save(path)

def create_pdf(path, text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    # Replacing utf-8 characters for fpdf simple usage or use generic ones
    text = text.replace("ü", "u").replace("ö", "o").replace("ı", "i").replace("ş", "s").replace("ğ", "g").replace("ç", "c")
    text = text.replace("Ü", "U").replace("Ö", "O").replace("İ", "I").replace("Ş", "S").replace("Ğ", "G").replace("Ç", "C")
    pdf.multi_cell(0, 10, txt=text)
    pdf.output(path)

def create_txt(path, text):
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)

files = [
    # Project: Luleburgaz
    ("IMG_4382.pdf", "Projeye ait Luleburgaz santiye beton dokum onayi ve test sonuclari. Sorumlu: Ali Yilmaz. Tarih: 12.08.2025.", create_pdf),
    ("Final_Hesap_2.xlsx", "Luleburgaz projesi icin hazirlanan hakedis ve kesin hesap tablosu. Guncellenmis versiyon.", create_xlsx),
    ("belediye_yeni.docx", "Luleburgaz Belediyesi imar plani basvurusu ve ruhsat yenileme talebi yazisi.", create_docx),
    ("luleburgaz_beton_test_sonuclari.pdf", "Projeye ait Luleburgaz santiye beton dokum onayi ve test sonuclari. Sorumlu: Ali Yilmaz. Tarih: 12.08.2025.", create_pdf), # exact duplicate

    # Project: Izmir Campus
    ("izmir_zemin.docx", "Izmir Kampus projesi zemin etudu raporu. Raporu hazirlayan zemin muhendisi.", create_docx),
    ("hakedis_rapor_izmir.xlsx", "Izmir kampus insaati 3. hakedis raporu. Odeme plani ve ilerleme.", create_xlsx),

    # Corporate
    ("personel_list.xlsx", "Dogrusal Muhendislik merkez ofis ve santiye personeli tam listesi ve maas bilgileri.", create_xlsx),
    ("Dogrusal_Sirket_Profili.pdf", "Dogrusal Muhendislik sirket vizyonu, tamamlanan projeler ve yonetim kurulu.", create_pdf),

    # Duplicate/Version check
    ("Final_Hesap_Son.xlsx", "Luleburgaz projesi icin hazirlanan kesin hesap tablosu. En son halidir.", create_xlsx),
    ("Final_Hesap_Son2.xlsx", "Luleburgaz projesi icin hazirlanan kesin hesap tablosu. Onaylandi bitti.", create_xlsx),

    # Orphan
    ("asdfgh.txt", "Sadece not alinmis anlamsiz yazi karalamalari.", create_txt)
]

for name, text, func in files:
    func(os.path.join("mock_archive", name), text)

print("Mock archive created.")
