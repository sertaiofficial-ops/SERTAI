import os
import PyPDF2
from docx import Document
from openpyxl import load_workbook

def parse_pdf(file_path):
    text = ""
    try:
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + " "
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
    return text.strip()

def parse_docx(file_path):
    text = ""
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            text += para.text + " "
    except Exception as e:
        print(f"Error reading DOCX {file_path}: {e}")
    return text.strip()

def parse_xlsx(file_path):
    text = ""
    try:
        # data_only=True ensures we read values, not formulas
        wb = load_workbook(file_path, data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        text += str(cell) + " "
    except Exception as e:
        print(f"Error reading XLSX {file_path}: {e}")
    return text.strip()

def parse_txt(file_path):
    text = ""
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
    except Exception as e:
        print(f"Error reading TXT {file_path}: {e}")
    return text.strip()

def extract_text_from_file(file_path):
    """
    Routes the file to the appropriate parser based on its extension.
    Returns the extracted text.
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        return parse_pdf(file_path)
    elif ext == '.docx':
        return parse_docx(file_path)
    elif ext == '.xlsx':
        return parse_xlsx(file_path)
    elif ext in ['.txt', '.md', '.csv']:
        return parse_txt(file_path)
    else:
        # Fallback or unsupported
        return ""
